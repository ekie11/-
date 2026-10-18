import gostdoc_engine.rules  # noqa: F401
from gostdoc_engine.analyze import analyze
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.rules.base import REGISTRY
from tests.builders import build_docx

KEY = "bibliography"
BOOK = "Иванов И. И. Название. — М. : Наука, 2020. — 100 с."


def check(blocks, pid="gost-7.32-2017"):
    cfg = builtin_profile(pid).rules[KEY]
    return REGISTRY[KEY]().check(analyze(build_docx(blocks)), cfg)


def test_no_section_no_refs():
    assert check([("p", "Просто текст.")]) == []


def test_no_section_with_ref():
    issues = check([("p", "Как показано в [1].")])
    assert len(issues) == 1 and issues[0].severity == "recommendation"
    assert issues[0].message == "Не найден раздел «Список использованных источников»"


def test_clean():
    assert check([("p", "Так [1], затем [2]."),
                  ("structural", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"),
                  ("source", f"1. {BOOK}"), ("source", f"2. {BOOK}")]) == []


def test_numbering_and_entry_problems():
    issues = check([("p", "Так [1] [2] [3] [4] [5]."),
                    ("structural", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"),
                    ("source", f"1. {BOOK}"),
                    ("source", "2. Иванов И. Статья // Журнал. — 2020."),
                    ("source", "4. Книга без года. — М. : Наука."),
                    ("source", "5. Сайт. — URL: http://x.ru, 2020. [Электронный ресурс]")])
    msgs = [i.message for i in issues]
    assert "Источник 4: нарушена нумерация, ожидался номер 3" in msgs
    assert "Источник 4: нет года издания" in msgs
    assert "Источник 5: у электронного ресурса нет даты обращения" in msgs
    assert "Источник 5: используйте «Текст : электронный» вместо «[Электронный ресурс]»" in msgs
    rec = [i for i in issues if i.severity == "recommendation"]
    assert len(rec) == 1
    assert all(not i.fixable for i in issues)


SEC = [("structural", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"),
       ("source", f"1. {BOOK}"), ("source", f"2. {BOOK}"), ("source", f"3. {BOOK}")]


def msgs(blocks, pid="gost-7.32-2017"):
    return [(i.message, i.severity) for i in check(blocks, pid)]


def test_find_citations():
    from gostdoc_engine.bibliography import find_citations
    m = analyze(build_docx([
        ("p", "См. [1], [2, с. 5], [1; 3], [1–3] и [2-3]."), ("p", "Нет."), *SEC]))
    assert [(i, n) for i, n in find_citations(m)] == [
        (0, [1, 2, 1, 3, 1, 2, 3, 2, 3])]
    big = analyze(build_docx([("p", "[1–500]"), *SEC]))
    assert len(find_citations(big)[0][1]) <= 50


def test_nonexistent_citation():
    r = msgs([("p", "Так [1] и [12]."), *SEC])
    assert ("Ссылка [12] указывает на несуществующий источник", "error") in r


def test_uncited_source_is_recommendation():
    r = msgs([("p", "Так [1][3]."), *SEC])
    assert ("На источник 2 нет ссылки в тексте", "recommendation") in r
    assert not any("На источник 1 " in m for m, _ in r)


def test_citation_in_toc_or_section_ignored():
    r = msgs([("p", "[1] [2] [3]"), *SEC[:1],
              ("source", f"1. {BOOK} См. [9]."), ("source", f"2. {BOOK}"),
              ("source", f"3. {BOOK}")])
    assert not any("[9]" in m for m, _ in r)


def test_mention_order():
    r = msgs([("p", "Сначала [2], потом [1] и [3]."), *SEC])
    order = [m for m, _ in r if m == "Источники пронумерованы не в порядке упоминания"]
    assert len(order) == 1
    assert msgs([("p", "Сначала [1], потом [2][1] и [3]."), *SEC]) == []


def test_fix_typography():
    import docx

    from gostdoc_engine.rules.base import REGISTRY as R
    d = build_docx(SEC[:1] + [
        ("source", "1. Москва-Наука, 2020. - М. :Наука,  2020. – 10 с. Текст:электронный"),
        ("source", "2. Статья//Журнал . - 2020. С. 10-15 и 10 - 15.")])
    model = analyze(d)
    cfg = builtin_profile("gost-7.32-2017").rules[KEY]
    assert R[KEY].fixable
    # разбить первый абзац источника на несколько прогонов
    par = d.paragraphs[1]
    t = par.runs[0].text
    par.runs[0].text = t[:22]
    par.add_run(t[22:40])
    par.add_run(t[40:])
    R[KEY]().fix(model, cfg)
    out = [p.text for p in d.paragraphs[1:]]
    assert out[0] == "1. Москва-Наука, 2020. — М. : Наука, 2020. — 10 с. Текст : электронный"
    assert out[1] == "2. Статья // Журнал . — 2020. С. 10-15 и 10 - 15."
    assert isinstance(d, docx.document.Document)


def test_fix_leaves_urls_and_emails():
    from gostdoc_engine.rules.base import REGISTRY as R
    src = ["1. URL: http://x.ru/a//b (дата обращения: 01.01.2020)",
           "2. e-mail:ivan@x.ru и mailto:a@b.ru"]
    d = build_docx(SEC[:1] + [("source", t) for t in src] + [("source", "3. Текст:электронный")])
    R[KEY]().fix(analyze(d), builtin_profile("gost-7.32-2017").rules[KEY])
    assert [p.text for p in d.paragraphs[1:]] == src[:2] + ["3. Текст : электронный"]


def test_citation_number_limits_and_roles():
    from gostdoc_engine.bibliography import find_citations
    m = analyze(build_docx([("p", "В [2019] и [0] году."),
                            ("table_caption", "Таблица 1 - Данные [1]"), *SEC]))
    assert find_citations(m) == []
