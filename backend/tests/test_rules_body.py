import pytest
from docx.oxml.ns import qn

import gostdoc_engine.rules  # noqa: F401
from gostdoc_engine.analyze import analyze
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.rules.base import REGISTRY
from tests.builders import build_docx

P = builtin_profile("gost-7.32-2017")
BLOCKS = [("p", "Обычный абзац основного текста."), ("p", "Второй абзац.")]


def check(key, doc):
    return REGISTRY[key]().check(analyze(doc), P.rules[key])


def fix(key, doc):
    REGISTRY[key]().fix(analyze(doc), P.rules[key])


KEYS = ["body.font_name", "body.font_size", "body.line_spacing",
        "body.first_line_indent", "body.alignment"]


@pytest.mark.parametrize("key", KEYS)
def test_correct_doc_clean(key):
    assert check(key, build_docx(BLOCKS)) == []


CASES = [
    ("body.font_name", {"font": "Arial"}, "Times New Roman", "Arial"),
    ("body.font_size", {"size_pt": 11}, "не менее 12 пт", "11 пт"),
    ("body.line_spacing", {"line_spacing": 1.0}, "1,5", "1"),
    ("body.first_line_indent", {"indent_cm": 0}, "1,25 см", "0 см"),
    ("body.alignment", {"align": "left"}, "по ширине", "по левому краю"),
]


@pytest.mark.parametrize(("key", "kw", "exp", "act"), CASES)
def test_violation_and_fix(key, kw, exp, act):
    doc = build_docx(BLOCKS, **kw)
    issues = check(key, doc)
    assert len(issues) == 2
    assert (issues[0].expected, issues[0].actual) == (exp, act)
    assert issues[0].paragraph_index == 0 and issues[0].rule == key
    fix(key, doc)
    assert check(key, doc) == []


def test_message_decimal_comma():
    msg = check("body.line_spacing", build_docx(BLOCKS, line_spacing=1.0))[0].message
    assert msg == "Межстрочный интервал: ожидается 1,5, найдено 1"


def test_size_fix_leaves_ok_paragraphs():
    doc = build_docx([("p", "Нормальный."), ("p", "Мелкий.")])
    for r in doc.paragraphs[1].runs:
        r.font.size = None
    doc.styles["Normal"].font.size = None
    # первый абзац: 14 в run; второй: наследует 8 от стиля
    from docx.shared import Pt
    doc.styles["Normal"].font.size = Pt(8)
    assert [i.paragraph_index for i in check("body.font_size", doc)] == [1]
    fix("body.font_size", doc)
    assert check("body.font_size", doc) == []
    assert doc.paragraphs[0].runs[0].font.size == Pt(14)


def test_run_level_font_detected_and_fixed():
    doc = build_docx(BLOCKS)
    doc.styles["Normal"].font.name = "Times New Roman"
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            rf = r._r.rPr.find(qn("w:rFonts"))
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), "Arial")
    assert len(check("body.font_name", doc)) == 2
    fix("body.font_name", doc)
    assert check("body.font_name", doc) == []
    assert all(r.font.name == "Times New Roman" for p in doc.paragraphs for r in p.runs)


def test_dominant_font_by_chars():
    from gostdoc_engine.units import effective_font_name
    doc = build_docx([("p", "x")])
    p = doc.paragraphs[0]
    p.runs[0].text = "ab"
    r2 = p.add_run("длинный текст")
    r2.font.name = "Arial"
    assert effective_font_name(p) == "Arial"


def test_table_cells_and_empty_paragraphs_ignored():
    doc = build_docx([("p", "Текст."), ("empty", ""), ("table", "a|b;c|d")], indent_cm=1.25)
    # ячейки: отступ 0, интервал 1.0, выравнивание left — не должно быть замечаний
    for key in KEYS:
        assert {i.paragraph_index for i in check(key, doc)} == set()
    doc2 = build_docx([("p", "Текст."), ("empty", ""), ("table", "a|b")],
                      indent_cm=0, line_spacing=1.0, align="left")
    idx = {i.paragraph_index for k in KEYS for i in check(k, doc2)}
    assert idx == {0}
    fix("body.first_line_indent", doc2)
    cell = doc2.tables[0].cell(0, 0).paragraphs[0]
    assert cell.paragraph_format.first_line_indent == 0


def test_image_only_paragraph_untouched():
    doc = build_docx([("figure", "alt"), ("p", "Текст.")], indent_cm=0)
    assert {i.paragraph_index for i in check("body.first_line_indent", doc)} == {1}


def test_list_item_skips_indent():
    doc = build_docx([("p", "Элемент.")], indent_cm=0)
    doc.paragraphs[0].style = doc.styles["List Bullet"]
    assert check("body.first_line_indent", doc) == []
    assert len(check("body.alignment", doc)) == 0


def test_style_level_properties_resolved():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    doc = build_docx([("p", "Текст.")])
    p = doc.paragraphs[0]
    pPr = p._p.pPr
    for tag in ("w:ind", "w:spacing", "w:jc"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    st = doc.styles["Normal"].paragraph_format
    st.line_spacing = 1.5
    st.first_line_indent = Cm(1.25)
    st.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for key in KEYS:
        assert check(key, doc) == []


def test_fix_does_not_modify_styles_or_table_cells():
    from docx.shared import Pt

    from gostdoc_engine.units import effective_font_name, effective_font_size_pt
    doc = build_docx([("p", "Текст."), ("table", "a|b")])
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(8)
    for r in doc.paragraphs[0].runs:
        r._r.rPr.remove(r._r.rPr.find(qn("w:rFonts")))
        r._r.rPr.remove(r._r.rPr.find(qn("w:sz")))
    cell = doc.tables[0].cell(0, 0).paragraphs[0]
    for r in cell.runs:
        r._r.rPr.remove(r._r.rPr.find(qn("w:rFonts")))
        r._r.rPr.remove(r._r.rPr.find(qn("w:sz")))
    body = doc.paragraphs[0]
    assert (effective_font_name(cell), effective_font_size_pt(cell)) == ("Arial", 8.0)
    fix("body.font_name", doc)
    fix("body.font_size", doc)
    assert check("body.font_name", doc) == [] and check("body.font_size", doc) == []
    assert effective_font_size_pt(body) == 14
    assert (effective_font_name(cell), effective_font_size_pt(cell)) == ("Arial", 8.0)
    assert (normal.font.name, normal.font.size) == ("Arial", Pt(8))
