from docx.enum.text import WD_ALIGN_PARAGRAPH

import gostdoc_engine.rules  # noqa: F401
from gostdoc_engine.analyze import analyze
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.rules.base import REGISTRY
from tests.builders import build_docx

P = builtin_profile("gost-7.32-2017")
KEY = "page_numbers"
BLOCKS = [("p", "Текст.")]


def check(doc):
    return REGISTRY[KEY]().check(analyze(doc), P.rules[KEY])


def fix(doc):
    REGISTRY[KEY]().fix(analyze(doc), P.rules[KEY])


def test_correct_clean():
    assert check(build_docx(BLOCKS)) == []


def test_no_page_numbers_added():
    doc = build_docx(BLOCKS, page_numbers=False)
    issues = check(doc)
    assert len(issues) == 1 and issues[0].fixable
    assert issues[0].paragraph_index is None
    assert issues[0].message == "Номер страницы: ожидается внизу страницы, найден отсутствует"
    fix(doc)
    assert check(doc) == []
    sec = doc.sections[0]
    assert sec.different_first_page_header_footer is True
    assert sec.footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "PAGE" in sec.footer._element.xml


def test_wrong_alignment():
    doc = build_docx(BLOCKS)
    doc.sections[0].footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    issues = check(doc)
    assert [(i.expected, i.actual) for i in issues] == [("по центру", "по правому краю")]
    fix(doc)
    assert check(doc) == []


def test_shown_on_first_page():
    doc = build_docx(BLOCKS)
    doc.sections[0].different_first_page_header_footer = False
    issues = check(doc)
    assert len(issues) == 1 and "первой странице" in issues[0].message
    fix(doc)
    assert check(doc) == []
    assert doc.sections[0].different_first_page_header_footer is True


def test_fix_idempotent_single_field():
    doc = build_docx(BLOCKS, page_numbers=False)
    fix(doc)
    fix(doc)
    assert doc.sections[0].footer._element.xml.count("PAGE") == 1


def test_numbers_text_footer_keeps_text():
    doc = build_docx(BLOCKS, page_numbers=False)
    doc.sections[0].footer.paragraphs[0].add_run("Колонтитул")
    fix(doc)
    assert "Колонтитул" in doc.sections[0].footer.paragraphs[0].text
    assert check(doc) == []


def _top_cfg():
    cfg = P.rules[KEY].model_copy(deep=True)
    cfg.value["position"] = "top_center"
    return cfg


def check_top(doc):
    return REGISTRY[KEY]().check(analyze(doc), _top_cfg())


def fix_top(doc):
    REGISTRY[KEY]().fix(analyze(doc), _top_cfg())


def test_top_footer_numbers_reported_and_moved():
    doc = build_docx(BLOCKS)  # номер внизу
    issues = check_top(doc)
    assert len(issues) == 1 and issues[0].fixable
    assert issues[0].message == "Номер страницы: ожидается вверху страницы, найден внизу страницы"
    fix_top(doc)
    assert check_top(doc) == []
    sec = doc.sections[0]
    assert "PAGE" in sec.header._element.xml
    assert "PAGE" not in sec.footer._element.xml
    assert sec.header.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert sec.different_first_page_header_footer is True


def test_top_missing():
    doc = build_docx(BLOCKS, page_numbers=False)
    issues = check_top(doc)
    assert [i.message for i in issues] == [
        "Номер страницы: ожидается вверху страницы, найден отсутствует"
    ]
    fix_top(doc)
    assert check_top(doc) == []


def test_top_idempotent_and_keeps_footer_text():
    doc = build_docx(BLOCKS)
    doc.sections[0].footer.paragraphs[0].add_run(" Колонтитул")
    fix_top(doc)
    xml1 = doc.sections[0].header._element.xml
    fix_top(doc)
    assert doc.sections[0].header._element.xml == xml1
    assert xml1.count("PAGE") == 1
    assert "Колонтитул" in doc.sections[0].footer.paragraphs[0].text


def test_top_wrong_alignment_and_first_page():
    doc = build_docx(BLOCKS, page_numbers=False)
    fix_top(doc)
    doc.sections[0].header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.sections[0].different_first_page_header_footer = False
    assert len(check_top(doc)) == 2
    fix_top(doc)
    assert check_top(doc) == []


def test_bottom_reports_header_numbers():
    doc = build_docx(BLOCKS)
    fix_top(doc)
    issues = check(doc)
    assert len(issues) == 1 and "вверху страницы" in issues[0].message
    fix(doc)
    assert check(doc) == []
    assert "PAGE" not in doc.sections[0].header._element.xml
