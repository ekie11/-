import io

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from gostdoc_api.guideline_text import GuidelineError, check_text, extract_text


def _pdf(pages: list[str]) -> bytes:
    w = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in pages:
        page = w.add_blank_page(width=200, height=200)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode() if text else b"")
        page[NameObject("/Contents")] = w._add_object(stream)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def test_docx_paragraph_and_table_in_order():
    d = Document()
    d.add_paragraph("Before table")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cell A"
    t.cell(0, 1).text = "Cell B"
    d.add_paragraph("After table")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text(buf.getvalue(), "m.DOCX")
    assert text.index("Before table") < text.index("Cell A") < text.index("Cell B")
    assert text.index("Cell B") < text.index("After table")


def test_pdf_page_markers():
    long = "Margin is 30 mm. " * 10
    text = extract_text(_pdf([long + "one", long + "two"]), "m.pdf")
    assert "[стр. 1]" in text and "[стр. 2]" in text
    assert text.index("[стр. 1]") < text.index("one") < text.index("[стр. 2]") < text.index("two")


def test_pdf_scan_rejected():
    with pytest.raises(GuidelineError) as e:
        extract_text(_pdf(["", ""]), "m.pdf")
    assert "похоже на скан" in str(e.value)
    assert e.value.status == 400


def test_unsupported_extension():
    with pytest.raises(GuidelineError) as e:
        extract_text(b"x", "m.txt")
    assert str(e.value) == "Поддерживаются DOCX и PDF"
    assert e.value.status == 400


def test_broken_files():
    for name in ("a.pdf", "a.docx"):
        with pytest.raises(GuidelineError) as e:
            extract_text(b"garbage", name)
        assert e.value.status == 400


def test_check_text_limits():
    assert check_text("abc") == "abc"
    with pytest.raises(GuidelineError) as e:
        check_text("  \n ")
    assert str(e.value) == "Текст методички пуст"
    with pytest.raises(GuidelineError) as e:
        check_text("a" * 150_001)
    assert e.value.status == 413
    assert "до 150 000 символов" in str(e.value)
    assert check_text("a" * 150_000)
