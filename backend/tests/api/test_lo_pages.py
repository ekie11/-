import io
import os
import shutil
import time
from pathlib import Path

import pytest
from docx import Document

from gostdoc_api import worker
from gostdoc_api.lo_pages import page_numbers_and_pdf
from gostdoc_api.pdf import PdfError
from gostdoc_engine.pipeline import process
from gostdoc_engine.profile import builtin_profile
from tests.builders import build_docx, save

from .test_jobs import DOCX_TYPE, fake_pdf, queue, redis_conn  # noqa: F401
from .test_profiles import GOST, reg

OK = """import sys, json
open(sys.argv[4], "wb").write(b"%PDF-1.4")
print(json.dumps({"_Toc00000001": 2, "_Toc00000002": 3}))
"""
GARBAGE = """import sys
open(sys.argv[4], "wb").write(b"%PDF-1.4")
print("not json")
"""
SLEEP = """import subprocess, sys, time
p = subprocess.Popen(["sleep", "30"])
open(sys.argv[3] + ".pid", "w").write(str(p.pid))
time.sleep(30)
"""
NO_PDF = 'print("{}")\n'


def _fake(tmp_path, monkeypatch, body, name="fake_lo.py"):
    path = tmp_path / name
    path.write_text(body, encoding="utf-8")
    monkeypatch.setenv("LO_UNO_SCRIPT", str(path))
    monkeypatch.setenv("LO_PYTHON", "python3")
    return path


def _docx(tmp_path):
    p = tmp_path / "in.docx"
    p.write_bytes(save(build_docx([("p", "Текст")])))
    return p


def test_success(tmp_path, monkeypatch):
    _fake(tmp_path, monkeypatch, OK)
    docx = _docx(tmp_path)
    before = docx.read_bytes()
    out = tmp_path / "o" / "r.pdf"
    pages = page_numbers_and_pdf(docx, out, soffice="x", timeout_s=10)
    assert pages == {"_Toc00000001": 2, "_Toc00000002": 3}
    assert out.read_bytes().startswith(b"%PDF")
    assert docx.read_bytes() == before


@pytest.mark.parametrize("body", [GARBAGE, NO_PDF, "import sys; sys.exit(3)\n"])
def test_bad_output(tmp_path, monkeypatch, body):
    _fake(tmp_path, monkeypatch, body)
    with pytest.raises(PdfError):
        page_numbers_and_pdf(_docx(tmp_path), tmp_path / "r.pdf", soffice="x", timeout_s=10)
    assert not (tmp_path / "r.pdf").exists()


def test_missing_interpreter(tmp_path, monkeypatch):
    _fake(tmp_path, monkeypatch, OK)
    monkeypatch.setenv("LO_PYTHON", str(tmp_path / "nope"))
    with pytest.raises(PdfError):
        page_numbers_and_pdf(_docx(tmp_path), tmp_path / "r.pdf", soffice="x", timeout_s=10)


def _zombie_children() -> list[str]:
    me = str(os.getpid())
    out = []
    for d in Path("/proc").iterdir():
        if d.name.isdigit():
            try:
                fields = (d / "stat").read_text().rsplit(")", 1)[1].split()
            except OSError:
                continue
            if fields[0] == "Z" and fields[1] == me:
                out.append(d.name)
    return out


def test_no_zombie_children_after_success_and_timeout(tmp_path, monkeypatch):
    _fake(tmp_path, monkeypatch, OK)
    page_numbers_and_pdf(_docx(tmp_path), tmp_path / "r.pdf", soffice="x", timeout_s=10)
    _fake(tmp_path, monkeypatch, SLEEP, "sleep_lo.py")
    with pytest.raises(PdfError):
        page_numbers_and_pdf(_docx(tmp_path), tmp_path / "r2.pdf", soffice="x", timeout_s=1)
    assert _zombie_children() == []


def test_timeout_kills_children(tmp_path, monkeypatch):
    _fake(tmp_path, monkeypatch, SLEEP)
    with pytest.raises(PdfError):
        page_numbers_and_pdf(_docx(tmp_path), tmp_path / "r.pdf", soffice="x", timeout_s=1)
    # pid-файл лежит во временном каталоге скрипта; ищем потомка по /proc
    deadline = time.monotonic() + 3
    while True:
        alive = []
        for d in Path("/proc").iterdir():
            if d.name.isdigit():
                try:
                    cmd = (d / "cmdline").read_bytes()
                except OSError:
                    continue
                if cmd.startswith(b"sleep\x0030"):
                    alive.append(d.name)
        if not alive:
            break
        assert time.monotonic() < deadline, f"остались процессы {alive}"
        time.sleep(0.05)


# ---------- воркер ----------

TOC_BLOCKS = [
    ("structural", "СОДЕРЖАНИЕ"),
    ("structural", "ВВЕДЕНИЕ"),
    ("p", "Текст введения."),
    ("h1", "1 Анализ"),
    ("p", "Основной текст раздела."),
    ("h1", "2 Выводы"),
    ("p", "Итог."),
]


def _toc_bytes(blocks=TOC_BLOCKS):
    p = builtin_profile("gost-7.32-2017")
    return process(save(build_docx(blocks)), p, "format").docx_bytes


def _toc_texts(doc):
    return [
        "".join(t.text or "" for t in link.iter("{*}t"))
        for link in doc.element.body.iter(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
        )
    ]


@pytest.mark.soffice
@pytest.mark.skipif(shutil.which("soffice") is None, reason="нет LibreOffice")
def test_real_libreoffice_page_numbers(tmp_path):
    from gostdoc_engine.toc import set_page_numbers

    src = tmp_path / "r.docx"
    src.write_bytes(_toc_bytes())
    t0 = time.monotonic()
    pages = page_numbers_and_pdf(src, tmp_path / "r.pdf", soffice="soffice", timeout_s=120)
    print("LO seconds", round(time.monotonic() - t0, 1))
    assert (tmp_path / "r.pdf").read_bytes().startswith(b"%PDF")
    assert len(pages) >= 3
    nums = [pages[k] for k in sorted(pages)]
    assert nums == sorted(nums) and len(set(nums)) == len(nums) and nums[0] > 1
    doc = Document(str(src))
    assert set_page_numbers(doc, pages) == len(pages)
    for text in _toc_texts(doc):
        assert "?" not in text


def _run_format(client, monkeypatch, tmp_path):
    reg(client)
    data = {"profile_id": GOST, "mode": "format"}
    r = client.post("/jobs", data=data, files={"file": ("a.docx", _toc_bytes_plain(), DOCX_TYPE)})
    return client.get(f"/jobs/{r.json()['id']}").json(), r.json()["id"]


def _toc_bytes_plain():
    return save(build_docx(TOC_BLOCKS))


@pytest.mark.usefixtures("queue")
def test_worker_fallback_sets_update_fields(client, monkeypatch, tmp_path):
    def boom(*a, **k):
        raise PdfError("x")

    monkeypatch.setattr(worker, "page_numbers_and_pdf", boom)
    got, jid = _run_format(client, monkeypatch, tmp_path)
    assert got["status"] == "done" and got["has_pdf"] is True  # старый render_pdf (фикстура)
    assert "Номера страниц в оглавлении обновятся при открытии в Word" in got["report"]["warnings"]
    body = client.get(f"/jobs/{jid}/docx").content
    doc = Document(io.BytesIO(body))
    uf = doc.settings.element.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields"
    )
    assert uf is not None


@pytest.mark.usefixtures("queue")
def test_worker_success_writes_numbers(client, monkeypatch, tmp_path):
    def fake(docx_path, out_pdf, *, soffice, timeout_s):
        doc = Document(str(docx_path))
        names = [
            lk.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor")
            for lk in doc.element.body.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
            )
        ]
        out_pdf.write_bytes(b"%PDF-1.4")
        return {n: i + 2 for i, n in enumerate(names)}

    monkeypatch.setattr(worker, "page_numbers_and_pdf", fake)
    got, jid = _run_format(client, monkeypatch, tmp_path)
    assert got["has_pdf"] is True
    assert "Номера страниц в оглавлении обновятся при открытии в Word" not in got["report"].get(
        "warnings", []
    )
    doc = Document(io.BytesIO(client.get(f"/jobs/{jid}/docx").content))
    texts = _toc_texts(doc)
    assert texts and all("?" not in t for t in texts)


@pytest.mark.usefixtures("queue")
def test_worker_no_toc_skips_lo(client, monkeypatch, tmp_path):
    def boom(*a, **k):
        raise AssertionError("не должен вызываться")

    monkeypatch.setattr(worker, "page_numbers_and_pdf", boom)
    reg(client)
    data = {"profile_id": GOST, "mode": "format"}
    content = save(build_docx([("p", "Просто текст")]))
    r = client.post("/jobs", data=data, files={"file": ("a.docx", content, DOCX_TYPE)})
    got = client.get(f"/jobs/{r.json()['id']}").json()
    assert got["status"] == "done"
    assert "Номера страниц в оглавлении обновятся при открытии в Word" not in got["report"].get(
        "warnings", []
    )


@pytest.mark.usefixtures("queue")
def test_worker_partial_pages_fall_back(client, monkeypatch, tmp_path):
    def fake(docx_path, out_pdf, *, soffice, timeout_s):
        out_pdf.write_bytes(b"%PDF-1.4")
        return {}

    monkeypatch.setattr(worker, "page_numbers_and_pdf", fake)
    got, _jid = _run_format(client, monkeypatch, tmp_path)
    assert "Номера страниц в оглавлении обновятся при открытии в Word" in got["report"]["warnings"]
