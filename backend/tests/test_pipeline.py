"""Тесты конвейера, защиты текста и CLI."""

import io
import subprocess
import sys

import docx
import pytest

from gostdoc_engine.analyze import analyze
from gostdoc_engine.importer import load_docx
from gostdoc_engine.pipeline import process
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.textguard import TextChangedError, normalized_words
from tests.builders import build_docx, save

P = builtin_profile("gost-7.32-2017")

BLOCKS = [
    ("structural", "введение"),
    ("p", "Текст «работы» — тест."),
    ("h1", "1 Анализ."),
    ("p", "Основной текст раздела."),
    ("h2", "1.1 Подраздел"),
    ("p", "Ещё текст."),
    ("figure", "Схема"),
    ("figure_caption", "Рисунок 1 - Схема"),
    ("table_caption", "Таблица 1 - Данные"),
    ("table", "a|b;c|d"),
    ("structural", "ЗАКЛЮЧЕНИЕ"),
    ("p", "Итог."),
]
BAD = dict(  # noqa: C408
    margins_mm=(10, 10, 10, 10),
    font="Arial",
    size_pt=10,
    line_spacing=1.0,
    indent_cm=0.0,
    align="left",
    use_heading_styles=False,
    page_numbers=False,
)


def _bad():
    return save(build_docx(BLOCKS, **BAD))


def _comments(data):
    return len(list(load_docx(data).comments))


def test_format_fixes_all_fixable():
    res = process(_bad(), P, "format")
    assert res.report["failed_rules"] == []
    assert [i for i in res.report["issues"] if i["fixable"]] == []


def test_format_is_idempotent():
    first = process(_bad(), P, "format")
    second = process(first.docx_bytes, P, "format")
    assert second.report["issues"] == []
    assert _comments(second.docx_bytes) == _comments(first.docx_bytes)


def test_reprocessing_does_not_pile_up_comments():
    first = process(_bad(), P, "format")
    second = process(first.docx_bytes, P, "format")
    assert _comments(second.docx_bytes) == _comments(first.docx_bytes)
    checked = process(second.docx_bytes, P, "check")
    assert _comments(checked.docx_bytes) == len(
        {i["paragraph_index"] for i in checked.report["issues"] if i["paragraph_index"] is not None}
        | ({None} if any(i["paragraph_index"] is None for i in checked.report["issues"]) else set())
    )


def test_repeated_check_keeps_comment_count():
    one = process(_bad(), P, "check")
    two = process(one.docx_bytes, P, "check")
    three = process(two.docx_bytes, P, "check")
    assert _comments(one.docx_bytes) > 0
    assert _comments(three.docx_bytes) == _comments(one.docx_bytes)


def test_user_comment_survives_format():
    d = build_docx(BLOCKS, **BAD)
    d.add_comment(list(d.paragraphs[1].runs), text="Моё", author="Иван", initials="И")
    out = load_docx(process(save(d), P, "format").docx_bytes)
    assert [c.author for c in out.comments if c.author == "Иван"] == ["Иван"]


def test_check_does_not_modify_formatting():
    res = process(_bad(), P, "check")
    assert res.report["issues"]
    sec = load_docx(res.docx_bytes).sections[0]
    assert round(sec.left_margin.mm) == 10
    assert load_docx(res.docx_bytes).styles["Normal"].font.name == "Arial"


@pytest.mark.parametrize("kw", [{}, BAD, {"use_heading_styles": False}, {"align": "left"}])
def test_text_preserved(kw):
    data = save(build_docx(BLOCKS, **kw))
    before = normalized_words(analyze(load_docx(data)))
    out = process(data, P, "format").docx_bytes
    assert normalized_words(analyze(load_docx(out))) == before


def test_normalized_words():
    d = docx.Document()
    d.add_paragraph("«А» – б в  „Г“")
    assert normalized_words(analyze(d)) == ['"а"', "-", "б", "в", '"г"']


def test_text_changed_error(monkeypatch):
    from gostdoc_engine import pipeline

    def evil(model, profile):
        model.paragraphs[0].para.runs[0].text = "другое"
        return []

    monkeypatch.setattr(pipeline, "run_fix", evil)
    with pytest.raises(TextChangedError):
        process(_bad(), P, "format")


def test_fix_failures_reported(monkeypatch):
    from gostdoc_engine import pipeline

    monkeypatch.setattr(pipeline, "run_fix", lambda m, p: ["page.size"])
    res = process(_bad(), P, "format")
    assert "page.size" in res.report["failed_rules"]
    assert len(res.report["failed_rules"]) == len(set(res.report["failed_rules"]))


def test_text_input():
    res = process("ВВЕДЕНИЕ\nТекст.\n1 Раздел\nТекст.", P, "format")
    assert res.report["mode"] == "format"
    assert res.report["profile"] == "gost-7.32-2017"
    assert set(res.report["summary"]) == {"errors", "recommendations"}


def test_custom_profile():
    data = P.model_dump()
    data["rules"]["body.font_size"]["value"].update(target_pt=12, min_pt=12)
    data["rules"]["page.margins"]["value"]["left_mm"] = 25
    data.update(id="custom", name="Свой", base="gost-7.32-2017", readonly=False)
    prof = type(P).model_validate(data)
    clean = [
        (
            k,
            t.replace("1 Анализ.", "1 Анализ")
            .replace("введение", "ВВЕДЕНИЕ")
            .replace(" - ", " — "),
        )
        for k, t in BLOCKS
    ]
    doc = save(build_docx(clean, size_pt=12, margins_mm=(20, 20, 25, 15)))
    res = process(doc, prof, "check")
    assert [i for i in res.report["issues"] if i["severity"] == "error"] == []


def test_cli_smoke(tmp_path):
    src = tmp_path / "in.txt"
    src.write_text("ВВЕДЕНИЕ\nТекст.\n", encoding="utf-8")
    out, rep = tmp_path / "out.docx", tmp_path / "r.json"
    r = subprocess.run(
        [
            sys.executable,
            "-m",
            "gostdoc_engine",
            str(src),
            "--text",
            "--profile",
            "gost-7.32-2017",
            "--mode",
            "format",
            "-o",
            str(out),
            "--report",
            str(rep),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert r.returncode == 0, r.stderr
    assert out.stat().st_size > 0 and rep.stat().st_size > 0
    docx.Document(io.BytesIO(out.read_bytes()))


def test_cli_input_error(tmp_path):
    src = tmp_path / "bad.docx"
    src.write_bytes(b"nope")
    r = subprocess.run(
        [
            sys.executable,
            "-m",
            "gostdoc_engine",
            str(src),
            "--profile",
            "gost-7.32-2017",
            "--mode",
            "check",
            "-o",
            str(tmp_path / "o.docx"),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert r.returncode == 2 and r.stderr.strip()


def _cli(*args):
    return subprocess.run(
        [sys.executable, "-m", "gostdoc_engine", *args],
        capture_output=True,
        text=True,
        check=False,
    )


def _assert_user_error(r):
    assert r.returncode == 2
    assert r.stderr.strip()
    assert "Traceback" not in r.stderr


def test_cli_unknown_profile(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("Текст.\n", encoding="utf-8")
    _assert_user_error(
        _cli(str(src), "--text", "--profile", "nope-1", "--mode", "check", "-o", str(tmp_path / "o.docx"))
    )


def test_cli_path_traversal_profile(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("Текст.\n", encoding="utf-8")
    _assert_user_error(
        _cli(str(src), "--text", "--profile", "../x", "--mode", "check", "-o", str(tmp_path / "o.docx"))
    )


def test_builtin_profile_rejects_bad_id():
    with pytest.raises(ValueError):
        builtin_profile("../x")
    with pytest.raises(ValueError):
        builtin_profile("nope-1")


def test_cli_missing_input(tmp_path):
    _assert_user_error(_cli(str(tmp_path / "none.docx"), "--mode", "check", "-o", str(tmp_path / "o.docx")))


def test_cli_non_utf8_text(tmp_path):
    src = tmp_path / "a.txt"
    src.write_bytes("Привет".encode("cp1251"))
    _assert_user_error(_cli(str(src), "--text", "--mode", "check", "-o", str(tmp_path / "o.docx")))


def test_body_final_period_drop_detected():
    d = docx.Document()
    d.add_paragraph("Обычный текст.")
    before = normalized_words(analyze(d))
    d.paragraphs[0].runs[0].text = "Обычный текст"
    assert normalized_words(analyze(d)) != before
