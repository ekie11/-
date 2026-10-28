import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH

import gostdoc_engine.rules  # noqa: F401
from gostdoc_engine.analyze import analyze
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.rules.base import REGISTRY
from tests.builders import build_docx

P = builtin_profile("gost-7.32-2017")
BLOCKS = [("structural", "ВВЕДЕНИЕ"), ("p", "Текст."), ("h1", "1 Анализ"),
          ("h2", "1.1 Подраздел"), ("p", "Текст.")]


def check(key, doc):
    return REGISTRY[key]().check(analyze(doc), P.rules[key])


def fix(key, doc):
    REGISTRY[key]().fix(analyze(doc), P.rules[key])


def msgs(key, doc):
    return [i.message for i in check(key, doc)]


@pytest.mark.parametrize("key", ["headings.section", "headings.structural"])
def test_correct_doc_clean(key):
    assert check(key, build_docx(BLOCKS)) == []


def test_section_bold():
    doc = build_docx(BLOCKS)
    for r in doc.paragraphs[2].runs:
        r.bold = False
    issues = check("headings.section", doc)
    assert [i.paragraph_index for i in issues] == [2]
    assert issues[0].message == "Заголовок: ожидается полужирное начертание, найдено обычное"
    assert issues[0].fixable
    fix("headings.section", doc)
    assert check("headings.section", doc) == []
    assert all(r.bold for r in doc.paragraphs[2].runs)


def test_section_align_indent_trailing_period():
    doc = build_docx(BLOCKS)
    p = doc.paragraphs[3]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.runs[0].text = "1.1 Подраздел."
    got = msgs("headings.section", doc)
    assert got == ["Заголовок: ожидается выравнивание по левому краю, найдено по центру",
                   "Заголовок: ожидается абзацный отступ 1,25 см, найдено 0 см",
                   "Заголовок: ожидается без точки в конце"]
    fix("headings.section", doc)
    assert check("headings.section", doc) == []
    assert p.text == "1.1 Подраздел"


def test_section_page_break_only_level1():
    doc = build_docx(BLOCKS)
    doc.paragraphs[2].paragraph_format.page_break_before = False
    doc.paragraphs[3].paragraph_format.page_break_before = True  # H2: безразлично
    issues = check("headings.section", doc)
    assert [i.paragraph_index for i in issues] == [2]
    assert "с новой страницы" in issues[0].message
    fix("headings.section", doc)
    assert check("headings.section", doc) == []
    assert doc.paragraphs[2].paragraph_format.page_break_before is True


def test_structural_uppercase_align_break():
    doc = build_docx([("structural", "Введение"), ("p", "Текст.")])
    p = doc.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.page_break_before = False
    got = msgs("headings.structural", doc)
    assert len(got) == 3
    assert "ожидается прописными буквами" in " ".join(got)
    fix("headings.structural", doc)
    assert check("headings.structural", doc) == []
    assert p.text == "ВВЕДЕНИЕ"
    assert doc.paragraphs[1].text == "Текст."


def test_style_bold_inherited_and_run_split():
    doc = build_docx([("h1", "1 Анализ")])
    p = doc.paragraphs[0]
    for r in p.runs:
        r.bold = None
    doc.styles["Heading 1"].font.bold = True
    assert check("headings.section", doc) == []
    doc.styles["Heading 1"].font.bold = False
    assert len(check("headings.section", doc)) == 1
    fix("headings.section", doc)
    assert check("headings.section", doc) == []
    assert doc.styles["Heading 1"].font.bold is False


def test_body_untouched():
    doc = build_docx(BLOCKS)
    fix("headings.section", doc)
    fix("headings.structural", doc)
    assert doc.paragraphs[1].text == "Текст."
    assert doc.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
