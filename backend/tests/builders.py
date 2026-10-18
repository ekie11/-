"""Генератор тестовых DOCX-документов. По умолчанию соответствует профилю ГОСТ 7.32-2017."""

import base64
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

Block = tuple[str, str]

_ALIGN = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
}

# Валидный PNG 1x1, генерируется из base64 в коде
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _set_font(run, font, size_pt, bold=None):
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def _fmt(p, align, indent_cm, line_spacing, page_break=False):
    pf = p.paragraph_format
    p.alignment = _ALIGN[align]
    pf.first_line_indent = Cm(indent_cm)
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if page_break:
        pf.page_break_before = True


def _add_page_number(section, font, size_pt):
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _set_font(run, font, size_pt)
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build_docx(
    blocks: list[Block],
    *,
    margins_mm=(20, 20, 30, 15),
    font="Times New Roman",
    size_pt=14,
    line_spacing=1.5,
    indent_cm=1.25,
    align="justify",
    use_heading_styles=True,
    page_numbers=True,
):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    top, bottom, left, right = margins_mm
    section.top_margin, section.bottom_margin = Mm(top), Mm(bottom)
    section.left_margin, section.right_margin = Mm(left), Mm(right)

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size_pt)

    def para(text, style, bold, a, indent, spacing=line_spacing, page_break=False):
        p = doc.add_paragraph(style=style)
        if text:
            _set_font(p.add_run(text), font, size_pt, bold)
        _fmt(p, a, indent, spacing, page_break)
        return p

    for kind, text in blocks:
        if kind == "structural":
            style = "Heading 1" if use_heading_styles else None
            para(text, style, True, "center", 0, page_break=True)
        elif kind in ("h1", "h2"):
            style = None
            if use_heading_styles:
                style = "Heading 1" if kind == "h1" else "Heading 2"
            para(text, style, True, "left", indent_cm, page_break=kind == "h1")
        elif kind == "p":
            para(text, None, None, align, indent_cm)
        elif kind == "figure":
            p = para("", None, None, "center", 0)
            p.add_run().add_picture(io.BytesIO(_PNG), width=Cm(1))
            p._p.xpath(".//wp:docPr")[0].set("descr", text)
        elif kind == "figure_caption":
            para(text, None, None, "center", 0)
        elif kind == "table_caption":
            para(text, None, None, "left", 0)
        elif kind == "source":
            para(text, None, None, "justify", indent_cm)
        elif kind == "empty":
            para("", None, None, align, indent_cm)
        elif kind == "table":
            rows = [r.split("|") for r in text.split(";")]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cp = table.cell(i, j).paragraphs[0]
                    _set_font(cp.add_run(val), font, size_pt)
                    _fmt(cp, "left", 0, 1.0)
        else:
            raise ValueError(f"Неизвестный вид блока: {kind}")

    if page_numbers:
        _add_page_number(section, font, size_pt)
    return doc


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
