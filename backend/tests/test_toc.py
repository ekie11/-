"""Тесты автоматического оглавления: стили заголовков, закладки, поле TOC."""

import io
import re

import docx
from docx.oxml.ns import qn

from gostdoc_engine import toc
from gostdoc_engine.analyze import analyze
from gostdoc_engine.importer import load_docx
from gostdoc_engine.pipeline import process
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.textguard import normalized_words
from tests.builders import build_docx, save

P = builtin_profile("gost-7.32-2017")

BLOCKS = [
    ("structural", "СОДЕРЖАНИЕ"),
    ("p", "Введение ........ 3"),
    ("p", "1 Анализ ........ 4"),
    ("structural", "ВВЕДЕНИЕ"),
    ("p", "Текст введения."),
    ("h1", "1 Анализ"),
    ("p", "Основной текст раздела."),
    ("h2", "1.1 Подраздел"),
    ("p", "Ещё текст."),
    ("structural", "ЗАКЛЮЧЕНИЕ"),
    ("p", "Итог."),
]


def _doc(blocks=BLOCKS, **kw):
    return build_docx(blocks, **kw)


def _reload(doc):
    return docx.Document(io.BytesIO(save(doc)))


def _instr(doc):
    return [t.text for t in doc.element.body.iter(qn("w:instrText"))]


def _links(doc):
    return list(doc.element.body.iter(qn("w:hyperlink")))


def _bookmarks(doc):
    return {b.get(qn("w:name")): b for b in doc.element.body.iter(qn("w:bookmarkStart"))}


def test_heading_styles_assigned_and_roles_unchanged():
    doc = _doc(use_heading_styles=False)
    model = analyze(doc)
    roles = [p.role for p in model.paragraphs]
    toc.apply_heading_styles(model)
    by_text = {p.text: p.para.style.name for p in model.paragraphs}
    assert by_text["1 Анализ"] == "Heading 1"
    assert by_text["1.1 Подраздел"] == "Heading 2"
    assert by_text["ВВЕДЕНИЕ"] == "Heading 1"
    assert [p.role for p in analyze(doc).paragraphs] == roles


def test_heading_look_preserved():
    doc = _doc(use_heading_styles=False)
    model = analyze(doc)
    h = next(p for p in model.paragraphs if p.text == "1 Анализ")
    before = (h.para.runs[0].font.size, h.para.paragraph_format.space_before)
    toc.apply_heading_styles(model)
    run = h.para.runs[0]
    assert (run.font.size, h.para.paragraph_format.space_before) == before
    assert run.font.color.rgb is not None and str(run.font.color.rgb) == "000000"


def test_build_toc_field_links_and_bookmarks():
    doc = _doc()
    model = analyze(doc)
    toc.apply_heading_styles(model)
    n = toc.build_toc(model)
    assert n == 4
    doc = _reload(doc)
    instr = _instr(doc)
    assert len(instr) == 1 and re.search(r'TOC \\o "1-3" \\h \\z \\u', instr[0])
    links = _links(doc)
    marks = _bookmarks(doc)
    assert len(links) == n
    for link in links:
        anchor = link.get(qn("w:anchor"))
        assert re.fullmatch(r"_Toc\d{8}", anchor)
        assert anchor in marks
    ids = [b.get(qn("w:id")) for b in marks.values()]
    assert len(set(ids)) == len(ids)

    # закладка стоит в абзаце своего заголовка
    def text(el):
        return "".join(t.text or "" for t in el.iter(qn("w:t")))

    texts = {link.get(qn("w:anchor")): text(link) for link in links}
    for name, el in marks.items():
        assert text(el.getparent()) in texts[name]


def test_toc_levels_styles_and_placeholder():
    doc = _doc()
    model = analyze(doc)
    toc.apply_heading_styles(model)
    toc.build_toc(model)
    doc = _reload(doc)
    model = analyze(doc)
    entries = [p for p in model.paragraphs if p.role.value == "TOC"]
    names = [p.para.style.name.lower() for p in entries]
    assert names == ["toc 1", "toc 1", "toc 2", "toc 1"]
    assert all(p.text.endswith("\t?") for p in entries)
    assert [p.text.split("\t")[0] for p in entries] == [
        "ВВЕДЕНИЕ",
        "1 Анализ",
        "1.1 Подраздел",
        "ЗАКЛЮЧЕНИЕ",
    ]
    style = next(s for s in doc.styles if (s.name or "").lower() == "toc 1")
    assert style.element.xpath('.//w:tab[@w:val="right"][@w:leader="dot"]')


def test_manual_toc_replaced():
    doc = _doc()
    model = analyze(doc)
    toc.apply_heading_styles(model)
    toc.build_toc(model)
    texts = [p.text for p in analyze(doc).paragraphs]
    assert "Введение ........ 3" not in texts and "1 Анализ ........ 4" not in texts


def test_idempotent_via_pipeline():
    bad = save(_doc(use_heading_styles=False))
    first = process(bad, P, "format")
    second = process(first.docx_bytes, P, "format")
    for res in (first, second):
        d = load_docx(res.docx_bytes)
        assert len(_instr(d)) == 1
        assert len(_links(d)) == 4
        assert len(_bookmarks(d)) == 4
    n1 = sorted(_bookmarks(load_docx(first.docx_bytes)))
    n2 = sorted(_bookmarks(load_docx(second.docx_bytes)))
    assert n1 == n2
    assert not [i for i in second.report["issues"] if i["rule"] == "toc"]


def test_words_outside_toc_unchanged():
    bad = save(_doc(use_heading_styles=False, size_pt=10))
    before = normalized_words(analyze(load_docx(bad)))
    res = process(bad, P, "format")  # TextChangedError, если слова изменились
    after = normalized_words(analyze(load_docx(res.docx_bytes)))
    assert after == before
    assert "введения." in after


def test_comments_survive():
    src = _doc(use_heading_styles=False)
    p = next(p for p in src.paragraphs if p.text == "Текст введения.")
    src.add_comment(p.runs, text="мой комментарий", author="Автор", initials="А")
    res = process(save(src), P, "format")
    d = load_docx(res.docx_bytes)
    assert any(c.author == "Автор" for c in d.comments)


def test_check_reports_manual_and_missing_toc():
    res = process(save(_doc()), P, "check")
    msgs = [i["message"] for i in res.report["issues"] if i["rule"] == "toc"]
    assert msgs == ["Содержание набрано вручную — будет заменено автоматическим оглавлением"]
    blocks = [("structural", "СОДЕРЖАНИЕ"), ("structural", "ВВЕДЕНИЕ"), ("p", "Текст.")]
    res = process(save(_doc(blocks)), P, "check")
    msgs = [i["message"] for i in res.report["issues"] if i["rule"] == "toc"]
    assert msgs == ["Нет автоматического оглавления"]
    issue = next(i for i in res.report["issues"] if i["rule"] == "toc")
    assert issue["severity"] == "recommendation" and issue["fixable"]


def test_no_toc_section_no_field():
    blocks = [("structural", "ВВЕДЕНИЕ"), ("p", "Текст."), ("h1", "1 Анализ")]
    res = process(save(_doc(blocks)), P, "format")
    d = load_docx(res.docx_bytes)
    assert _instr(d) == [] and not [i for i in res.report["issues"] if i["rule"] == "toc"]


def test_set_page_numbers_and_update_fields():
    doc = _doc()
    model = analyze(doc)
    toc.apply_heading_styles(model)
    toc.build_toc(model)
    links = _links(doc)
    pages = {links[0].get(qn("w:anchor")): 3, links[1].get(qn("w:anchor")): 12}
    toc.set_page_numbers(doc, pages)
    texts = ["".join(t.text for t in link.iter(qn("w:t"))) for link in _links(doc)]
    assert texts[0].endswith("3") and texts[1].endswith("12") and texts[2].endswith("?")
    toc.set_update_fields(doc)
    toc.set_update_fields(doc)
    d = _reload(doc)
    assert len(d.settings.element.findall(qn("w:updateFields"))) == 1
    assert d.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"


def test_toc_runs_after_headings_regardless_of_profile_order():
    from gostdoc_engine.profile import Profile

    data = P.model_dump()
    rules = data["rules"]
    reordered = {"toc": rules.pop("toc"), **rules}
    prof = Profile.model_validate({**data, "rules": reordered})
    res = process(save(_doc(use_heading_styles=False)), prof, "format")
    assert res.report["failed_rules"] == []
    assert len(_instr(load_docx(res.docx_bytes))) == 1


def test_profiles_have_toc():
    for pid in ("gost-7.32-2017", "gost-2.105-2019", "espd-19.106-78"):
        cfg = builtin_profile(pid).rules["toc"]
        assert cfg.value == {"levels": 3, "hyperlinks": True}
        assert cfg.severity == "recommendation"


T = [("p", "Министерство"), ("p", "Отчёт"), ("structural", "СОДЕРЖАНИЕ")]


def _format_words(blocks):
    bad = save(build_docx(blocks, use_heading_styles=False))
    before = [t for t in (p.text for p in load_docx(bad).paragraphs) if t]
    res = process(bad, P, "format")
    d = load_docx(res.docx_bytes)
    return before, [p.text for p in d.paragraphs], d


def test_repro_a_pages_no_loss():
    blocks = T + [
        ("p", "1 Обзор ....... 4"),
        ("h1", "1 Обзор"),
        ("p", "Важный текст."),
        ("p", "Ещё абзац."),
    ]
    _, after, d = _format_words(blocks)
    assert "Важный текст." in after and "Ещё абзац." in after
    assert "1 Обзор" in after and len(_instr(d)) == 1 and len(_links(d)) == 1
    assert "1 Обзор ....... 4" not in after


def test_repro_b_entries_without_pages():
    blocks = T + [
        ("p", "Введение"),
        ("p", "1 Обзор"),
        ("structural", "ВВЕДЕНИЕ"),
        ("p", "Текст."),
        ("h1", "1 Обзор"),
        ("p", "Тело."),
    ]
    _, after, d = _format_words(blocks)
    assert "Тело." in after and "Текст." in after
    assert len(_instr(d)) == 1 and len(_links(d)) == 2 and len(_bookmarks(d)) == 2
    assert "Введение" not in after


def test_repro_c_struct_after():
    blocks = T + [
        ("p", "1 Обзор ....... 4"),
        ("p", "Заключение ..... 9"),
        ("h1", "1 Обзор"),
        ("p", "Тело."),
        ("structural", "ЗАКЛЮЧЕНИЕ"),
        ("p", "Вывод."),
    ]
    _, after, d = _format_words(blocks)
    assert "Тело." in after and "1 Обзор" in after and "Вывод." in after
    assert "Заключение ..... 9" not in after
    assert len(_links(d)) == 2


def test_e_body_line_with_number_stays():
    blocks = T + [
        ("p", "Выручка за 2020 год составила  15"),
        ("p", "Дальше текст."),
        ("h1", "1 Обзор"),
        ("p", "Тело."),
    ]
    _, after, d = _format_words(blocks)
    assert "Выручка за 2020 год составила  15" in after and len(_links(d)) == 0


def test_f_sentence_equal_to_later_heading_in_region_is_entry():
    # допущение: строка сразу после СОДЕРЖАНИЕ, равная заголовку ниже, считается пунктом
    blocks = T + [("p", "1 Обзор"), ("p", "Дальше текст."), ("h1", "1 Обзор"), ("p", "Тело.")]
    _, after, d = _format_words(blocks)
    assert "Дальше текст." in after and "Тело." in after
    assert len(_links(d)) == 1 and "1 Обзор" not in [t for t in after if "\t" not in t][:3]


def test_deleting_body_paragraph_still_raises(monkeypatch):
    import pytest

    from gostdoc_engine.rules.body import FontNameRule
    from gostdoc_engine.textguard import TextChangedError

    def evil(self, model, cfg):
        for p in model.paragraphs:
            if p.text == "Тело.":
                p.para._p.getparent().remove(p.para._p)

    monkeypatch.setattr(FontNameRule, "fix", evil)
    blocks = T + [("structural", "ВВЕДЕНИЕ"), ("h1", "1 Обзор"), ("p", "Тело.")]
    with pytest.raises(TextChangedError):
        process(save(build_docx(blocks, use_heading_styles=False)), P, "format")


def test_check_silent_without_headings():
    res = process(save(build_docx(T + [("p", "Введение ....... 3")])), P, "check")
    assert not [i for i in res.report["issues"] if i["rule"] == "toc"]


MISMATCH = (
    "Содержание набрано вручную и не совпадает с заголовками — автоматическое "
    "оглавление не вставлено, проверьте содержание вручную"
)


def _skipped(blocks):
    before, after, d = _format_words(blocks)
    assert [t.casefold() for t in after] == [t.casefold() for t in before]
    assert _instr(d) == [] and _links(d) == []
    res = process(save(build_docx(blocks, use_heading_styles=False)), P, "format")
    msgs = [i["message"] for i in res.report["issues"] if i["rule"] == "toc"]
    assert msgs == [MISMATCH]


def test_unmatched_leader_lines_skip_with_issue():
    # a: «Введение ... 3» не совпадает ни с каким заголовком
    _skipped(
        T
        + [
            ("p", "Введение ....... 3"),
            ("p", "1 Обзор ....... 4"),
            ("h1", "1 Обзор"),
            ("p", "Важный текст."),
        ]
    )
    # c
    _skipped(
        T
        + [
            ("p", "Введение ....... 3"),
            ("p", "1 Обзор ....... 4"),
            ("h1", "1 Обзор"),
            ("p", "Тело."),
            ("structural", "ЗАКЛЮЧЕНИЕ"),
            ("p", "Вывод."),
        ]
    )


def test_original_b_no_pages_skipped():
    _skipped(
        T
        + [
            ("p", "Введение"),
            ("p", "1 Обзор"),
            ("h1", "1 Обзор"),
            ("p", "Важный текст тела раздела."),
        ]
    )


def test_leader_line_is_never_heading():
    from gostdoc_engine.analyze import classify_text

    assert classify_text("1 Обзор ....... 4", True) is None
    assert classify_text("1 Обзор\t4", True) is None
