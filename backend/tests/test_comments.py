import io
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from gostdoc_engine.analyze import analyze
from gostdoc_engine.comments import add_comments
from gostdoc_engine.rules.base import Issue
from tests.builders import build_docx, save

BLOCKS = [("p", "Первый абзац."), ("p", "Второй абзац.")]


def issue(source, message, paragraph_index):
    return Issue(
        rule="body.font_name",
        source=source,
        severity="error",
        message=message,
        expected="",
        actual="",
        paragraph_index=paragraph_index,
        fixable=True,
    )


def reload_comments(doc):
    """Сохранить, перечитать пакет и вернуть w:comment из word/comments.xml."""
    with zipfile.ZipFile(io.BytesIO(save(doc))) as zf:
        if "word/comments.xml" not in zf.namelist():
            return []
        root = etree.fromstring(zf.read("word/comments.xml"))
    return root.findall(qn("w:comment"))


def comment_text(comment):
    return "\n".join(
        "".join(t.text or "" for t in p.iter(qn("w:t"))) for p in comment.findall(qn("w:p"))
    )


def test_comments_written_with_author_and_source():
    doc = build_docx(BLOCKS)
    model = analyze(doc)
    issues = [issue("п. 6.1.1", "Шрифт: ожидается Times New Roman", 0)]

    assert add_comments(model, issues) == 1

    comments = reload_comments(doc)
    assert len(comments) == 1
    assert comments[0].get(qn("w:author")) == "ГОСТ-сервис"
    assert comments[0].get(qn("w:initials")) == "ГС"
    assert "[п. 6.1.1] Шрифт: ожидается Times New Roman" in comment_text(comments[0])


def test_issues_on_same_paragraph_are_merged():
    doc = build_docx(BLOCKS)
    model = analyze(doc)
    issues = [
        issue("п. 6.1.1", "Первое", 0),
        issue("п. 6.2.1", "Второе", 0),
        issue("п. 6.1.1", "Третье", 1),
    ]

    assert add_comments(model, issues) == 2

    texts = [comment_text(c) for c in reload_comments(doc)]
    assert texts == [
        "[п. 6.1.1] Первое\n[п. 6.2.1] Второе",
        "[п. 6.1.1] Третье",
    ]


def test_page_issue_goes_to_first_non_empty_paragraph():
    doc = build_docx([("empty", ""), ("p", "Текст абзаца.")])
    model = analyze(doc)
    issues = [issue("п. 6.1.1", "Поля: 25 мм", None)]

    assert add_comments(model, issues) == 1

    comments = reload_comments(doc)
    assert "Поля: 25 мм" in comment_text(comments[0])
    anchored = doc.element.body.xpath(".//w:commentRangeStart")[0].getparent()
    assert [t.text for t in anchored.iter(qn("w:t"))] == ["Текст абзаца."]


def test_paragraph_without_runs_gets_anchor_run():
    doc = build_docx([("empty", ""), ("p", "Текст.")])
    model = analyze(doc)
    empty_index = model.paragraphs[0].index
    assert model.paragraphs[0].para.runs == []

    assert add_comments(model, [issue("п. 6.1.1", "Замечание", empty_index)]) == 1

    assert len(reload_comments(doc)) == 1
    anchored = doc.element.body.xpath(".//w:commentRangeStart")[0].getparent()
    assert anchored.find(qn("w:r")) is not None
    assert [t.text for t in anchored.iter(qn("w:t"))] == []


def test_no_issues_adds_nothing():
    doc = Document()
    assert add_comments(analyze(doc), []) == 0
    assert reload_comments(doc) == []


def test_remove_service_comments_keeps_foreign():
    from gostdoc_engine.comments import remove_service_comments

    doc = build_docx(BLOCKS)
    model = analyze(doc)
    add_comments(model, [issue("п. 1", "Наше", 0)])
    runs = list(model.paragraphs[1].para.runs)
    doc.add_comment(runs, text="Чужое", author="Иван", initials="И")

    assert remove_service_comments(doc) == 1

    comments = reload_comments(doc)
    assert [comment_text(c) for c in comments] == ["Чужое"]
    body = doc.element.body
    assert len(body.xpath(".//w:commentRangeStart")) == 1
    assert len(body.xpath(".//w:commentRangeEnd")) == 1
    assert len(body.xpath(".//w:commentReference")) == 1
    assert remove_service_comments(build_docx(BLOCKS)) == 0
