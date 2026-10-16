import io

from docx import Document
from docx.shared import Cm, Mm

from tests.builders import build_docx, save


def reopen(doc):
    return Document(io.BytesIO(save(doc)))


def test_defaults_margins_and_indent():
    doc = reopen(build_docx([("structural", "ВВЕДЕНИЕ"), ("p", "Текст.")]))
    s = doc.sections[0]
    assert abs(s.top_margin.mm - 20) < 0.1
    assert abs(s.bottom_margin.mm - 20) < 0.1
    assert abs(s.left_margin.mm - 30) < 0.1
    assert abs(s.right_margin.mm - 15) < 0.1
    p = next(p for p in doc.paragraphs if p.text == "Текст.")
    assert abs(p.paragraph_format.first_line_indent.cm - 1.25) < 0.01
    assert p.paragraph_format.line_spacing == 1.5
    assert p.runs[0].font.size.pt == 14
    assert p.runs[0].font.name == "Times New Roman"


def test_blocks_kinds():
    doc = reopen(
        build_docx(
            [
                ("h1", "1 Обзор"),
                ("h2", "1.1 Подраздел"),
                ("figure", "alt"),
                ("figure_caption", "Рисунок 1 — Схема"),
                ("table_caption", "Таблица 1 — Данные"),
                ("table", "a|b;c|d"),
                ("empty", ""),
            ]
        )
    )
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert doc.paragraphs[1].style.name == "Heading 2"
    assert doc.paragraphs[2]._p.xpath(".//w:drawing")
    assert doc.paragraphs[3].text.startswith("Рисунок 1")
    assert doc.paragraphs[3].alignment == 1
    assert doc.paragraphs[4].paragraph_format.first_line_indent in (0, Cm(0))
    assert len(doc.tables) == 1 and doc.tables[0].cell(1, 0).text == "c"
    # caption directly above the table
    body = list(doc.element.body)
    assert body.index(doc.tables[0]._tbl) == body.index(doc.paragraphs[4]._p) + 1


def test_structural_heading_format():
    doc = reopen(build_docx([("structural", "ВВЕДЕНИЕ")]))
    p = doc.paragraphs[0]
    assert p.alignment == 1 and p.runs[0].bold
    assert p.paragraph_format.page_break_before


def test_no_heading_styles_and_violations():
    doc = reopen(
        build_docx(
            [("h1", "1 Обзор"), ("p", "Текст.")],
            use_heading_styles=False,
            font="Arial",
            size_pt=12,
            line_spacing=1.0,
            indent_cm=0,
            align="left",
            margins_mm=(10, 10, 10, 10),
        )
    )
    assert doc.paragraphs[0].style.name == "Normal" and doc.paragraphs[0].runs[0].bold
    p = doc.paragraphs[1]
    assert p.runs[0].font.name == "Arial" and p.runs[0].font.size.pt == 12
    assert p.alignment == 0 and p.paragraph_format.line_spacing == 1.0
    assert abs(doc.sections[0].left_margin.mm - Mm(10).mm) < 0.1


def test_page_numbers():
    doc = reopen(build_docx([("p", "x")]))
    s = doc.sections[0]
    assert s.different_first_page_header_footer
    assert "PAGE" in s.footer._element.xml
    assert s.footer.paragraphs[0].alignment == 1
    doc2 = reopen(build_docx([("p", "x")], page_numbers=False))
    assert "PAGE" not in doc2.sections[0].footer._element.xml
