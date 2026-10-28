from docx.enum.text import WD_ALIGN_PARAGRAPH

import gostdoc_engine.rules  # noqa: F401
from gostdoc_engine.analyze import analyze
from gostdoc_engine.profile import builtin_profile
from gostdoc_engine.rules.base import REGISTRY
from tests.builders import build_docx

P = builtin_profile("gost-7.32-2017")
OK = [
    ("p", "Текст."),
    ("figure", "alt"),
    ("figure_caption", "Рисунок 1 — Схема"),
    ("table_caption", "Таблица 1 — Данные"),
    ("table", "a|b;c|d"),
    ("p", "Конец."),
]


def check(key, doc):
    return REGISTRY[key]().check(analyze(doc), P.rules[key])


def fix(key, doc):
    REGISTRY[key]().fix(analyze(doc), P.rules[key])


def test_correct_clean():
    doc = build_docx(OK)
    assert check("captions.figure", doc) == []
    assert check("captions.table", doc) == []


def test_caption_wrong_dash_fixed():
    for dash in ("-", "–"):
        doc = build_docx([("figure", "a"), ("figure_caption", f"Рисунок 1 {dash} Схема")])
        issues = check("captions.figure", doc)
        assert len(issues) == 1 and issues[0].fixable
        assert issues[0].message == f"Подпись рисунка: ожидается тире «—», найдено «{dash}»"
        fix("captions.figure", doc)
        assert doc.paragraphs[1].text == "Рисунок 1 — Схема"
        assert check("captions.figure", doc) == []


def test_dash_fix_across_runs():
    doc = build_docx([("table_caption", "Таблица 2 - Данные"), ("table", "a|b")])
    p = doc.paragraphs[0]
    p.runs[0].text = "Таблица 2 "
    p.add_run("- Данные")
    fix("captions.table", doc)
    assert p.text == "Таблица 2 — Данные"


def test_pattern_mismatch_reported_not_fixable():
    doc = build_docx([("figure", "a"), ("figure_caption", "Рисунок 1 Схема")])
    issues = check("captions.figure", doc)
    assert len(issues) == 1 and not issues[0].fixable
    fix("captions.figure", doc)
    assert doc.paragraphs[1].text == "Рисунок 1 Схема"


def test_alignment_fixed():
    doc = build_docx(OK)
    doc.paragraphs[2].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.paragraphs[3].alignment = WD_ALIGN_PARAGRAPH.CENTER
    assert len(check("captions.figure", doc)) == 1
    assert len(check("captions.table", doc)) == 1
    fix("captions.figure", doc)
    fix("captions.table", doc)
    assert check("captions.figure", doc) == [] and check("captions.table", doc) == []


def test_figure_caption_above_reported_not_moved():
    doc = build_docx([("figure_caption", "Рисунок 1 — Схема"), ("figure", "a"), ("p", "Т.")])
    issues = check("captions.figure", doc)
    assert len(issues) == 1 and not issues[0].fixable
    assert "под рисунком" in issues[0].message
    fix("captions.figure", doc)
    assert doc.paragraphs[0].text == "Рисунок 1 — Схема"


def test_figure_caption_after_empty_paragraph_ok():
    doc = build_docx([("figure", "a"), ("empty", ""), ("figure_caption", "Рисунок 1 — Схема")])
    assert check("captions.figure", doc) == []


def test_table_caption_after_table_reported_not_moved():
    doc = build_docx([("table", "a|b"), ("table_caption", "Таблица 1 — Данные")])
    issues = check("captions.table", doc)
    assert len(issues) == 1 and not issues[0].fixable
    assert "над таблицей" in issues[0].message
    fix("captions.table", doc)
    body = doc.element.body
    assert body[0].tag.endswith("}tbl")
    assert doc.paragraphs[0].text == "Таблица 1 — Данные"
    assert len(check("captions.table", doc)) == 1


def test_table_caption_indent_fixed():
    doc = build_docx(OK)
    doc.paragraphs[3].paragraph_format.first_line_indent = None
    from docx.shared import Cm

    doc.paragraphs[3].paragraph_format.first_line_indent = Cm(1.25)
    assert len(check("captions.table", doc)) == 1
    fix("captions.table", doc)
    assert check("captions.table", doc) == []
