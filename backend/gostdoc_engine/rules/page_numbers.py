"""Правило нумерации страниц: поле PAGE в нижнем или верхнем колонтитуле (position)."""

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..model import DocModel
from ..profile import RuleConfig
from ..units import _style_chain
from .base import Issue, Rule, register

_PAGE_FIELD = re.compile(r"^\s*PAGE\b", re.IGNORECASE)
_ALIGN_NAMES = {
    "justify": "по ширине",
    "left": "по левому краю",
    "center": "по центру",
    "right": "по правому краю",
}


def _has_page_field(par) -> bool:
    if any(_PAGE_FIELD.match(t.text or "") for t in par._p.xpath(".//w:instrText")):
        return True
    return any(
        _PAGE_FIELD.match(f.get(qn("w:instr")) or "") for f in par._p.xpath(".//w:fldSimple")
    )


def _alignment(par) -> str:
    """Выравнивание абзаца колонтитула (части стилей документа у колонтитула нет)."""
    ppr = par._p.pPr
    chain = [ppr] + [st.element.pPr for st in _style_chain(par)]
    for pr in chain:
        jc = pr.find(qn("w:jc")) if pr is not None else None
        if jc is not None:
            return {"both": "justify", "center": "center", "right": "right", "end": "right"}.get(
                jc.get(qn("w:val")), "left"
            )
    return "left"


def _field_paragraph(part):
    return next((p for p in part.paragraphs if _has_page_field(p)), None)


POSITIONS = {
    "bottom_center": ("footer", "header", "внизу страницы", "вверху страницы"),
    "top_center": ("header", "footer", "вверху страницы", "внизу страницы"),
}


def _position(cfg: RuleConfig) -> str:
    pos = cfg.value.get("position", "bottom_center")
    return pos if pos in POSITIONS else "bottom_center"


def _own_sections(model: DocModel, attr: str):
    """Разделы с собственным колонтитулом нужного вида (первый раздел — всегда)."""
    for i, s in enumerate(model.doc.sections):
        if i == 0 or not getattr(s, attr).is_linked_to_previous:
            yield i, s


def _existing_field_paragraph(section, attr: str):
    part = getattr(section, attr)
    # у связанного колонтитула своего содержимого нет (и обращаться к нему не надо)
    return None if part.is_linked_to_previous else _field_paragraph(part)


def _remove_page_fields(par) -> None:
    """Убирает из абзаца только поля PAGE (остальной текст не трогает)."""
    for el in par._p.xpath(".//w:fldSimple"):
        if _PAGE_FIELD.match(el.get(qn("w:instr")) or ""):
            el.getparent().remove(el)
    group: list = []
    depth = 0
    for r in list(par._p.xpath("./w:r")):
        chars = [c.get(qn("w:fldCharType")) for c in r.findall(qn("w:fldChar"))]
        if depth == 0 and "begin" not in chars:
            continue
        group.append(r)
        depth += chars.count("begin") - chars.count("end")
        if depth <= 0:
            depth = 0
            instr = "".join(t.text or "" for g in group for t in g.findall(qn("w:instrText")))
            if _PAGE_FIELD.match(instr):
                for g in group:
                    g.getparent().remove(g)
            group = []


@register
class PageNumbersRule(Rule):
    key = "page_numbers"
    fixable = True

    def _issue(self, cfg, msg, exp, act):
        return Issue(
            rule=self.key,
            source=cfg.source,
            severity=cfg.severity,
            message=msg,
            expected=exp,
            actual=act,
            paragraph_index=None,
            fixable=True,
        )

    def check(self, model: DocModel, cfg: RuleConfig) -> list[Issue]:
        attr, other, place, other_place = POSITIONS[_position(cfg)]
        issues = []
        for i, s in _own_sections(model, attr):
            where = f" (раздел {i + 1})" if len(model.doc.sections) > 1 else ""
            par = _field_paragraph(getattr(s, attr))
            stray = _existing_field_paragraph(s, other)
            if par is None:
                act = other_place if stray is not None else "отсутствует"
                issues.append(
                    self._issue(
                        cfg, f"Номер страницы{where}: ожидается {place}, найден {act}", place, act
                    )
                )
            else:
                act = _alignment(par)
                if act != "center":
                    e, a = _ALIGN_NAMES["center"], _ALIGN_NAMES[act]
                    issues.append(
                        self._issue(
                            cfg,
                            f"Номер страницы{where}: ожидается выравнивание {e}, найдено {a}",
                            e,
                            a,
                        )
                    )
                if stray is not None:
                    issues.append(
                        self._issue(
                            cfg,
                            f"Номер страницы{where}: ожидается только {place}, найден также "
                            f"{other_place}",
                            f"только {place}",
                            f"также {other_place}",
                        )
                    )
            if (
                par is not None
                and i == 0
                and not cfg.value.get("show_on_first", False)
                and not s.different_first_page_header_footer
            ):
                issues.append(
                    self._issue(
                        cfg,
                        "Номер на первой странице: ожидается скрыт, найден показан",
                        "скрыт",
                        "показан",
                    )
                )
        return issues

    def fix(self, model: DocModel, cfg: RuleConfig) -> None:
        attr, other, _, _ = POSITIONS[_position(cfg)]
        for i, s in _own_sections(model, attr):
            part = getattr(s, attr)
            if i == 0 and part.is_linked_to_previous:
                part.is_linked_to_previous = False
            par = _field_paragraph(part)
            if par is None:
                par = part.paragraphs[0]
                run = par.add_run()
                for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
                    if kind:
                        el = OxmlElement("w:fldChar")
                        el.set(qn("w:fldCharType"), kind)
                    else:
                        el = OxmlElement("w:instrText")
                        el.set(qn("xml:space"), "preserve")
                        el.text = text
                    run._r.append(el)
            if _alignment(par) != "center":
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0 and not cfg.value.get("show_on_first", False):
                s.different_first_page_header_footer = True
            stray = _existing_field_paragraph(s, other)
            while stray is not None:
                _remove_page_fields(stray)
                stray = _existing_field_paragraph(s, other)
