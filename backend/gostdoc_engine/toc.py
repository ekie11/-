"""Автоматическое оглавление: стили заголовков, закладки _Toc, поле TOC с готовым результатом."""

import re

from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.text.paragraph import Paragraph

from . import units
from .analyze import _style_heading_role, looks_like_leader_line, toc_field_paragraphs, toc_keys
from .model import DocModel, DocParagraph, Role

TOC_INSTR = ' TOC \\o "1-3" \\h \\z \\u '
PLACEHOLDER = "?"
_BOOKMARK = re.compile(r"_Toc\d+")
_BOOKMARK8 = re.compile(r"_Toc(\d{8})")
_FIRST_NUMBER = 10000001
_LEVEL_STYLE = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}
_HEADING_LEVEL = {Role.HEADING_1: 1, Role.HEADING_2: 2, Role.HEADING_3: 3}
_DEFAULT_WIDTH_TWIPS = 9355


def _norm(text: str) -> str:
    return " ".join(text.split()).upper()


def toc_heading(model: DocModel) -> DocParagraph | None:
    """Абзац «СОДЕРЖАНИЕ» (структурный элемент)."""
    return next(
        (p for p in model.paragraphs if p.role == Role.STRUCTURAL and _norm(p.text) == "СОДЕРЖАНИЕ"),
        None,
    )


def has_toc_field(doc) -> bool:
    body = doc.element.body
    for el in body.iter(qn("w:instrText")):
        if re.match(r"\s*TOC\b", el.text or ""):
            return True
    return any(re.match(r"\s*TOC\b", e.get(qn("w:instr")) or "")
               for e in body.iter(qn("w:fldSimple")))


def _entry_level(p: DocParagraph) -> int | None:
    if p.role == Role.STRUCTURAL:
        return None if _norm(p.text) == "СОДЕРЖАНИЕ" else 1
    return _HEADING_LEVEL.get(p.role)


def _entries(model: DocModel) -> list[tuple[DocParagraph, int]]:
    out = []
    for p in model.paragraphs:
        level = _entry_level(p)
        if level is not None and p.text.strip():
            out.append((p, level))
    return out


_HEADLIKE = (Role.STRUCTURAL, Role.HEADING_1, Role.HEADING_2, Role.HEADING_3)


def manual_toc_problem(model: DocModel) -> bool:
    """Ручное содержание не удаётся надёжно сопоставить с заголовками — оглавление не строится."""
    heading = toc_heading(model)
    if heading is None:
        return False
    after = [p for p in model.paragraphs
             if p.index > heading.index and p.role not in (Role.EMPTY, Role.IN_TABLE)]
    styled = {p.index: _style_heading_role(p.para) is not None for p in after}
    # строки с отточием/страницей, не подтверждённые заголовками, до следующего раздела
    for p in after:
        if p.role == Role.STRUCTURAL or styled[p.index]:
            break
        if p.role != Role.TOC and looks_like_leader_line(p.text):
            return True
    if not after or after[0].role == Role.TOC or styled[after[0].index]:
        return False
    first = after[0]
    if first.role not in _HEADLIKE:
        return True  # обычный текст сразу после СОДЕРЖАНИЕ — не ясно, что это
    # два неоформленных заголовка подряд — вероятно, список без номеров страниц
    return (len(after) > 1 and after[1].role in _HEADLIKE and not styled[after[1].index])


def has_entries(model: DocModel) -> bool:
    return bool(_entries(model))


# ---------- стили заголовков ----------

def _heading_style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        pass
    style = doc.styles.add_style(name, 1)
    level = int(name[-1])
    style.base_style = doc.styles["Normal"]
    style.next_paragraph_style = doc.styles["Normal"]
    ppr = style.element.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    ppr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level - 1}"/>'))
    return style


def _old_element(par: Paragraph, kind: str, tag: str):
    """Первый элемент tag в цепочке стилей, затем в умолчаниях документа."""
    for style in units._style_chain(par):
        props = style.element.pPr if kind == "pPr" else style.element.rPr
        if props is not None and props.find(qn(tag)) is not None:
            return props.find(qn(tag))
    defaults = units._doc_defaults(par, kind)
    return defaults.find(qn(tag)) if defaults is not None else None


def _pin_look(par: Paragraph) -> None:
    """Закрепить прямым форматированием то, что абзац наследовал от прежнего стиля."""
    ppr = par._p.get_or_add_pPr()
    for tag, getter in (("w:spacing", ppr.get_or_add_spacing), ("w:ind", ppr.get_or_add_ind),
                        ("w:jc", ppr.get_or_add_jc)):
        old = _old_element(par, "pPr", tag)
        direct = ppr.find(qn(tag))
        if old is None or direct is not None:
            continue
        el = getter()
        for k, v in old.attrib.items():
            el.set(k, v)
    for run in units.all_runs(par):
        rpr = run._r.get_or_add_rPr()
        for tag, getter in (("w:rFonts", rpr.get_or_add_rFonts), ("w:sz", rpr.get_or_add_sz),
                            ("w:b", rpr.get_or_add_b), ("w:i", rpr.get_or_add_i),
                            ("w:color", rpr.get_or_add_color)):
            if rpr.find(qn(tag)) is not None:
                continue
            old = _old_element(par, "rPr", tag)
            if old is None and tag == "w:color":
                # у нового стиля цвета нет — оставлять нечего; у прежнего тоже
                continue
            if old is None:
                continue
            el = getter()
            for k, v in old.attrib.items():
                el.set(k, v)


def apply_heading_styles(model: DocModel) -> int:
    """Назначить стили Heading 1/2/3; внешний вид прежний (закрепляется прямым форматированием)."""
    changed = 0
    for p, level in _entries(model):
        name = _LEVEL_STYLE[level]
        if p.para.style is not None and p.para.style.name == name:
            continue
        style = _heading_style(model.doc, name)
        _pin_look(p.para)
        _neutralize_style_look(p.para, style)
        p.para.style = style
        changed += 1
    return changed


def _neutralize_style_look(par: Paragraph, style) -> None:
    """Свойства нового стиля, которых не было у старого (цвет, подчёркивание), отменяются явно."""
    new = style.element.rPr
    if new is None:
        return
    for run in units.all_runs(par):
        rpr = run._r.get_or_add_rPr()
        if new.find(qn("w:color")) is not None and rpr.find(qn("w:color")) is None:
            rpr.get_or_add_color().set(qn("w:val"), "000000")
        if new.find(qn("w:u")) is not None and rpr.find(qn("w:u")) is None:
            rpr.get_or_add_u().set(qn("w:val"), "none")
        for tag in ("w:b", "w:i"):
            if new.find(qn(tag)) is not None and rpr.find(qn(tag)) is None:
                rpr.append(parse_xml(f'<{tag} {nsdecls("w")} w:val="0"/>'))


# ---------- закладки ----------

def _bookmark_state(doc):
    names, max_id = set(), 0
    for b in doc.element.body.iter(qn("w:bookmarkStart")):
        names.add(b.get(qn("w:name")))
        try:
            max_id = max(max_id, int(b.get(qn("w:id"))))
        except (TypeError, ValueError):
            pass
    nums = [int(m.group(1)) for n in names if n and (m := _BOOKMARK8.fullmatch(n))]
    return max_id, max(nums, default=_FIRST_NUMBER - 1)


def _existing_bookmark(par: Paragraph) -> str | None:
    for b in par._p.findall(qn("w:bookmarkStart")):
        name = b.get(qn("w:name")) or ""
        if _BOOKMARK.fullmatch(name):
            return name
    return None


def _add_bookmark(par: Paragraph, name: str, bid: int) -> None:
    p = par._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    ppr = p.pPr
    if ppr is not None:
        ppr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)


# ---------- стили оглавления ----------

def _find_style_id(doc, level: int) -> str | None:
    want = f"toc {level}"
    for s in doc.styles.element.findall(qn("w:style")):
        name = s.find(qn("w:name"))
        if name is not None and (name.get(qn("w:val")) or "").lower() == want:
            return s.get(qn("w:styleId"))
    return None


def _text_width_twips(doc) -> int:
    sec = doc.sections[0]
    try:
        return int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
    except TypeError:
        return _DEFAULT_WIDTH_TWIPS


def _body_font(model: DocModel) -> tuple[str | None, float | None]:
    body = next((p for p in model.paragraphs if p.role == Role.BODY), None)
    if body is None:
        return None, None
    return units.effective_font_name(body.para), units.effective_font_size_pt(body.para)


def _ensure_toc_style(model: DocModel, level: int) -> str:
    doc = model.doc
    found = _find_style_id(doc, level)
    if found:
        return found
    normal = doc.styles["Normal"].style_id
    font, size = _body_font(model)
    rpr = ""
    if font or size:
        fonts = (f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/>'
                 if font else "")
        sz = f'<w:sz w:val="{int(size * 2)}"/><w:szCs w:val="{int(size * 2)}"/>' if size else ""
        rpr = f"<w:rPr>{fonts}{sz}</w:rPr>"
    left = f'<w:ind w:left="{220 * (level - 1)}"/>' if level > 1 else ""
    xml = (
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="TOC{level}">'
        f'<w:name w:val="toc {level}"/><w:basedOn w:val="{normal}"/><w:next w:val="{normal}"/>'
        '<w:autoRedefine/><w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
        f'<w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{_text_width_twips(doc)}"/></w:tabs>'
        f'<w:spacing w:before="0" w:after="0"/>{left}</w:pPr>{rpr}</w:style>'
    )
    doc.styles.element.append(parse_xml(xml))
    return f"TOC{level}"


# ---------- построение поля ----------

def _fld_run(kind: str):
    r = OxmlElement("w:r")
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), kind)
    r.append(f)
    return r


def _instr_run():
    r = OxmlElement("w:r")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = TOC_INSTR
    r.append(t)
    return r


def _text_run(text: str):
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _entry_paragraph(style_id: str, anchor: str, text: str, first: bool, last: bool):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), style_id)
    ppr.append(ps)
    p.append(ppr)
    if first:
        p.append(_fld_run("begin"))
        p.append(_instr_run())
        p.append(_fld_run("separate"))
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    link.append(_text_run(text))
    tab = OxmlElement("w:r")
    tab.append(OxmlElement("w:tab"))
    link.append(tab)
    link.append(_text_run(PLACEHOLDER))
    p.append(link)
    if last:
        p.append(_fld_run("end"))
    return p


def _is_toc_sdt(el) -> bool:
    if el.tag != qn("w:sdt"):
        return False
    gallery = el.xpath(".//w:docPartGallery/@w:val")
    if any("table of contents" in g.lower() for g in gallery):
        return True
    return any(re.match(r"\s*TOC\b", t.text or "") for t in el.iter(qn("w:instrText")))


def _verify_old_toc(model: DocModel, after: list[DocParagraph]) -> None:
    """Страховка: удаляются только пункты оглавления (в поле или повторённые ниже как заголовки)."""
    in_field = toc_field_paragraphs(model.doc)
    heads = [q for q in model.paragraphs
             if q.role in (Role.STRUCTURAL, Role.HEADING_1, Role.HEADING_2, Role.HEADING_3)]
    for p in after:
        if p.role != Role.TOC or p.para._p in in_field:
            continue
        keys = toc_keys(p.text)
        if not any(q.index > p.index and toc_keys(q.text) & keys for q in heads):
            raise RuntimeError("Не удалось заменить содержание автоматически")


def _remove_old_toc(model: DocModel, heading: DocParagraph) -> None:
    after = [p for p in model.paragraphs if p.index > heading.index]
    _verify_old_toc(model, after)
    stop = next((p for p in after if p.role == Role.STRUCTURAL), None)
    stop_el = stop.para._p if stop is not None else None
    for sib in list(heading.para._p.itersiblings()):
        if sib is stop_el:
            break
        if _is_toc_sdt(sib):
            sib.getparent().remove(sib)
    for p in after:
        if p.role == Role.TOC:
            el = p.para._p
            if el.getparent() is not None:
                el.getparent().remove(el)


def build_toc(model: DocModel) -> int:
    """Вставить (или заменить) поле оглавления после «СОДЕРЖАНИЕ»; вернуть число пунктов."""
    heading = toc_heading(model)
    entries = _entries(model)
    if heading is None or not entries or manual_toc_problem(model):
        return 0
    doc = model.doc
    _verify_old_toc(model, [p for p in model.paragraphs if p.index > heading.index])
    max_id, max_num = _bookmark_state(doc)
    anchors = []
    for p, _level in entries:
        name = _existing_bookmark(p.para)
        if name is None:
            max_id += 1
            max_num += 1
            name = f"_Toc{max_num:08d}"
            _add_bookmark(p.para, name, max_id)
        anchors.append(name)
    style_ids = {lv: _ensure_toc_style(model, lv) for lv in {lv for _, lv in entries}}
    _remove_old_toc(model, heading)
    anchor_el = heading.para._p
    for i, ((p, level), name) in enumerate(zip(entries, anchors, strict=True)):
        text = " ".join(p.text.split())
        el = _entry_paragraph(style_ids[level], name, text, i == 0, i == len(entries) - 1)
        anchor_el.addnext(el)
        anchor_el = el
    return len(entries)


# ---------- номера страниц и обновление полей ----------

def set_page_numbers(doc, pages: dict[str, int]) -> int:
    """Вписать номера страниц в ссылки оглавления (ключ — имя закладки); вернуть число обновлённых."""
    done = 0
    for link in doc.element.body.iter(qn("w:hyperlink")):
        anchor = link.get(qn("w:anchor"))
        if anchor not in pages:
            continue
        ts = link.xpath("./w:r/w:t")
        if len(ts) < 2:
            continue
        ts[-1].text = str(pages[anchor])
        done += 1
    return done


_AFTER_UPDATE_FIELDS = (
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
    "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
    "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
    "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol",
    "listSeparator",
)


def set_update_fields(doc) -> None:
    """Попросить Word обновить поля при открытии (w:updateFields в settings.xml)."""
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is not None:
        existing.set(qn("w:val"), "true")
        return
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    for child in settings:
        if child.tag in {qn(f"w:{n}") for n in _AFTER_UPDATE_FIELDS}:
            child.addprevious(el)
            return
    settings.append(el)


__all__ = ["apply_heading_styles", "build_toc", "has_toc_field", "set_page_numbers",
           "set_update_fields", "toc_heading"]
