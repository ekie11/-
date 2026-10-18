"""Загрузка DOCX и импорт обычного текста с понятными ошибками."""

import io
import re

import docx

from .analyze import classify_text
from .model import Role

MAX_BYTES = 20 * 1024 * 1024
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"
_ENCRYPTED = "EncryptedPackage".encode("utf-16-le")
_NOT_DOCX = "Файл не похож на документ Word (.docx)"
_STYLES = {
    Role.STRUCTURAL: "Heading 1",
    Role.HEADING_1: "Heading 1",
    Role.HEADING_2: "Heading 2",
    Role.HEADING_3: "Heading 3",
}


class InputError(Exception):
    """Ошибка входных данных; сообщение предназначено пользователю."""


def load_docx(data: bytes):
    if len(data) > MAX_BYTES:
        raise InputError("Файл больше 20 МБ")
    if data.startswith(_OLE_MAGIC):
        if _ENCRYPTED in data:
            raise InputError("Документ защищён паролем — снимите защиту и загрузите снова")
        raise InputError(_NOT_DOCX)
    try:
        return docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise InputError(_NOT_DOCX) from exc


def text_to_docx(text: str):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]*\t[ \t]*", " ", line)
        line = line.strip()
        if line:
            lines.append(line)
    if not lines:
        raise InputError("Текст пустой — вставьте текст работы")
    doc = docx.Document()
    for line in lines:
        role = classify_text(line, bold=False)
        doc.add_paragraph(line, style=_STYLES.get(role, "Normal"))
    return doc
