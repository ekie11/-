"""Замечания как комментарии Word: один комментарий на абзац."""

from docx.text.run import Run

from .model import DocModel, DocParagraph
from .rules.base import Issue

AUTHOR = "ГОСТ-сервис"
INITIALS = "ГС"


def _first_non_empty(model: DocModel) -> DocParagraph | None:
    return next((p for p in model.paragraphs if p.text.strip()), None)


def add_comments(model: DocModel, issues: list[Issue]) -> int:
    """Добавить комментарии Word к абзацам; вернуть их число.

    Замечания одного абзаца объединяются в один комментарий. Замечания без абзаца
    (например, page.*) относятся к первому непустому абзацу. Замечания к неизвестному
    абзацу пропускаются.
    """
    by_index = {p.index: p for p in model.paragraphs}
    first = _first_non_empty(model)
    grouped: dict[int, list[str]] = {}
    for item in issues:
        if item.paragraph_index is None:
            target = first
        else:
            target = by_index.get(item.paragraph_index)
        if target is None:
            continue
        grouped.setdefault(target.index, []).append(f"[{item.source}] {item.message}")

    for index, lines in grouped.items():
        par = by_index[index].para
        runs: list[Run] = list(par.runs)
        if not runs:
            runs = [par.add_run("")]
        model.doc.add_comment(runs, text="\n".join(lines), author=AUTHOR, initials=INITIALS)
    return len(grouped)
